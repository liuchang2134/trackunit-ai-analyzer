from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

root=Path.cwd()
source=root/'docs/机联智检_完整参赛Proposal.md'
out=root/'docs/机联智检_完整参赛Proposal.docx'
doc=Document()
sec=doc.sections[0]
sec.page_width=Inches(8.5); sec.page_height=Inches(11)
sec.top_margin=Inches(.7); sec.bottom_margin=Inches(.65)
sec.left_margin=Inches(.78); sec.right_margin=Inches(.78)
sec.header_distance=Inches(.28); sec.footer_distance=Inches(.28)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
    style=doc.styles[name]
    style.font.name='Calibri'
    style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei' if name!='Normal' else 'SimSun')
    style.font.color.rgb=RGBColor.from_string('000000')
lang=OxmlElement('w:lang');lang.set(qn('w:val'),'zh-CN');lang.set(qn('w:eastAsia'),'zh-CN');doc.styles['Normal'].element.get_or_add_rPr().append(lang)
normal=doc.styles['Normal']; normal.font.size=Pt(11)
normal.paragraph_format.line_spacing=1.18
normal.paragraph_format.space_after=Pt(7)
for name,size in [('Title',25),('Heading 1',17),('Heading 2',12.5)]:
    st=doc.styles[name]; st.font.size=Pt(size); st.font.bold=True
    st.paragraph_format.space_before=Pt(8); st.paragraph_format.space_after=Pt(9)
    st.paragraph_format.keep_with_next=True
for style in doc.styles:
    for border in list(style.element.iter(qn('w:pBdr'))): border.getparent().remove(border)
    if style.type == 1:
        pp=style.element.get_or_add_pPr()
        for tag in ['kinsoku','overflowPunct']:
            el=OxmlElement('w:'+tag);el.set(qn('w:val'),'1');pp.append(el)
h=sec.header.paragraphs[0]; h.text='机联智检  |  完整参赛项目计划书'; h.style='Normal'
for r in h.runs:r.font.size=Pt(8);r.font.color.rgb=RGBColor.from_string('000000')
f=sec.footer.paragraphs[0]; f.alignment=WD_ALIGN_PARAGRAPH.RIGHT
f.add_run('2026年9月14日  ·  ')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
for r in f.runs:r.font.size=Pt(8)

def rich(p,text):
    for part in re.split(r'(\*\*.*?\*\*|\[[^\]]+\]\([^\)]+\))',text):
        if not part:continue
        if part.startswith('**'):
            p.add_run(part[2:-2]).bold=True
        elif re.match(r'\[[^\]]+\]\(',part):
            m=re.match(r'\[([^\]]+)\]\(([^\)]+)\)',part)
            link=OxmlElement('w:hyperlink');link.set(qn('r:id'),p.part.relate_to(m[2],'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True))
            r=OxmlElement('w:r');rp=OxmlElement('w:rPr');c=OxmlElement('w:color');c.set(qn('w:val'),'185A8D');rp.append(c);r.append(rp);t=OxmlElement('w:t');t.text=m[1];r.append(t);link.append(r);p._p.append(link)
        else:p.add_run(re.sub(r'([，。；：！？])', lambda m: chr(0x2060)+m[1], part.replace('`','')))

def table(lines):
    data=[[c.strip() for c in line.strip('|').split('|')] for line in lines if not re.match(r'^\|[\s:|\-]+\|$',line)]
    t=doc.add_table(rows=0, cols=len(data[0]));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    widths=[1.38,2.64,2.92] if len(data[0])==3 else [1.8,5.14]
    for col,w in zip(t.columns,widths):col.width=Inches(w)
    pr=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    pr.append(borders)
    for idx,row in enumerate(data):
        cells=t.add_row().cells
        trpr=t.rows[-1]._tr.get_or_add_trPr();lock=OxmlElement('w:cantSplit');trpr.append(lock)
        if idx==0:
            repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
        for j,(cell,txt) in enumerate(zip(cells,row)):
            cell.width=Inches(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp=cell._tc.get_or_add_tcPr();marg=OxmlElement('w:tcMar')
            for side in ['top','bottom','left','right']:
                el=OxmlElement('w:'+side);el.set(qn('w:w'),'90');el.set(qn('w:type'),'dxa');marg.append(el)
            cp.append(marg)
            sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E8EFF5' if idx==0 else ('F7F9FB' if idx%2==0 else 'FFFFFF'));cp.append(sh)
            p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(1);p.paragraph_format.space_before=Pt(1);p.paragraph_format.line_spacing=1.12
            p.paragraph_format.keep_with_next=idx==0
            rich(p,txt)
            for r in p.runs:r.font.size=Pt(10);r.bold=idx==0
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=Pt(4);p.add_run().font.size=Pt(4)

lines=source.read_text(encoding='utf-8-sig').splitlines();i=0;next_page=False
while i<len(lines):
    line=lines[i].strip()
    if not line:i+=1;continue
    if line=='<!-- pagebreak -->':next_page=True
    elif line.startswith('|'):
        block=[]
        while i<len(lines) and lines[i].strip().startswith('|'):block.append(lines[i].strip());i+=1
        table(block);continue
    elif line.startswith('# '):rich(doc.add_paragraph(style='Title'),line[2:])
    elif line.startswith('## '):
        p=doc.add_paragraph(style='Heading 1');p.paragraph_format.page_break_before=next_page;next_page=False;rich(p,line[3:].replace('、',' '))
    elif line.startswith('### '):rich(doc.add_paragraph(style='Heading 2'),line[4:])
    else:rich(doc.add_paragraph(),line)
    i+=1
props=doc.core_properties;props.title='机联智检完整参赛项目计划书';props.subject='面向既有车联网平台的本地AI诊断与备件辅助工具';props.author='';props.keywords='Proposal;工程机械;车联网;本地AI;备件映射'
doc.save(out)
print(out)
print('Characters:',len(source.read_text(encoding='utf-8-sig')))

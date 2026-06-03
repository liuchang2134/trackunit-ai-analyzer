from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.data_store import ROOT, cache_status, load_machines
from app.trend_analysis import get_fleet_trends


REPORT_DIR = ROOT / "data" / "reports"


def _ensure_report_dir() -> None:
    REPORT_DIR.mkdir(parents=True, exist_ok=True)


def _timestamp() -> str:
    return datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")


def build_fleet_excel_report(days: int = 30) -> Path:
    _ensure_report_dir()
    trends = get_fleet_trends(days)
    cache = cache_status()
    machines = load_machines()
    output = REPORT_DIR / f"fleet_report_{_timestamp()}.xlsx"

    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Summary"
    summary_sheet.append(["Metric", "Value"])
    summary_sheet.append(["Generated At", datetime.now(timezone.utc).isoformat()])
    summary_sheet.append(["Data Source", cache.get("data_source")])
    summary_sheet.append(["Cache Updated At", cache.get("updated_at")])
    summary_sheet.append(["Cache Fresh", cache.get("fresh")])
    if trends["latest"]:
        for key, value in trends["latest"].items():
            summary_sheet.append([key, value])

    machine_sheet = workbook.create_sheet("Machines")
    machine_sheet.append(["Machine ID", "Serial Number", "Model", "Type", "Customer", "Location", "Last Seen"])
    for machine in machines:
        machine_sheet.append([
            machine.machine_id,
            machine.serial_number,
            machine.model,
            machine.machine_type,
            machine.customer,
            machine.location,
            machine.last_seen_at,
        ])

    trend_sheet = workbook.create_sheet("Fleet Trends")
    trend_sheet.append([
        "Captured At",
        "Total",
        "Online",
        "Offline",
        "Low Fuel",
        "Low Utilization",
        "Fault Records",
        "Average Fuel %",
        "Average Operating Hours",
        "Source",
    ])
    for point in trends["points"]:
        trend_sheet.append([
            point.get("captured_at"),
            point.get("total_machines"),
            point.get("online_machines"),
            point.get("offline_machines"),
            point.get("low_fuel_machines"),
            point.get("low_utilization_machines"),
            point.get("fault_records"),
            point.get("average_fuel_percent"),
            point.get("average_operating_hours"),
            point.get("source"),
        ])

    workbook.save(output)
    return output


def build_fleet_pdf_report(days: int = 30) -> Path:
    _ensure_report_dir()
    trends = get_fleet_trends(days)
    cache = cache_status()
    output = REPORT_DIR / f"fleet_report_{_timestamp()}.pdf"
    styles = getSampleStyleSheet()
    story = [
        Paragraph("Trackunit Fleet AI Analyzer Report", styles["Title"]),
        Paragraph(f"Generated at: {datetime.now(timezone.utc).isoformat()}", styles["Normal"]),
        Paragraph(f"Data source: {cache.get('data_source')} | Cache fresh: {cache.get('fresh')}", styles["Normal"]),
        Spacer(1, 12),
    ]

    if trends["latest"]:
        latest = trends["latest"]
        summary_data = [
            ["Metric", "Value"],
            ["Total Machines", latest.get("total_machines")],
            ["Online Machines", latest.get("online_machines")],
            ["Offline Machines", latest.get("offline_machines")],
            ["Low Fuel Machines", latest.get("low_fuel_machines")],
            ["Low Utilization Machines", latest.get("low_utilization_machines")],
            ["Fault Records", latest.get("fault_records")],
        ]
    else:
        summary_data = [["Metric", "Value"], ["History", "No fleet history available yet"]]

    table = Table(summary_data, hAlign="LEFT")
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0f4fd8")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#d5deea")),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("PADDING", (0, 0), (-1, -1), 6),
    ]))
    story.extend([table, Spacer(1, 12)])

    story.append(Paragraph("Trend Notes", styles["Heading2"]))
    for note in trends["summary"]:
        story.append(Paragraph(f"- {note}", styles["Normal"]))

    doc = SimpleDocTemplate(str(output), pagesize=letter)
    doc.build(story)
    return output
